"""로컬 PDF 편집과 OCR. 원본은 읽기만 하고 결과를 새 이름으로 저장한다."""
from __future__ import annotations

import csv
from contextlib import contextmanager
import ctypes
from ctypes import wintypes
from dataclasses import dataclass
from functools import wraps
import io
import os
from pathlib import Path
import re
import shutil
import signal
import statistics
import subprocess
import sys
import tempfile
import threading
import time

from PIL import Image, ImageOps, ImageSequence
from pypdf import PdfReader, PdfWriter
import pypdfium2 as pdfium
from pdf_lock import PDFIUM_LOCK


def _serialized_pdfium(function):
    """PDF 도구와 미리보기가 같은 네이티브 런타임에 동시에 진입하지 않는다."""
    @wraps(function)
    def guarded(*args, **kwargs):
        with PDFIUM_LOCK:
            return function(*args, **kwargs)
    return guarded


class PdfToolError(RuntimeError):
    def __init__(self, code, detail="", *, source=""):
        self.code, self.detail, self.source = code, str(detail), str(source)
        super().__init__(code + (": " + self.detail if self.detail else ""))


class Cancelled(PdfToolError):
    def __init__(self):
        super().__init__("cancelled")


@contextmanager
def source_context(source):
    """실제로 처리 중인 입력을 예외에 붙인다. 표시 문자열에서 경로를 추정하지 않는다."""
    try:
        yield
    except Cancelled:
        raise
    except PdfToolError as error:
        if not error.source:
            error.source = str(source)
        raise
    except Exception as error:
        raise PdfToolError("processing_failed", str(error), source=str(source)) from error


@dataclass
class Request:
    operation: str
    paths: list[str]
    pages: str = ""
    rotation: int = 90
    compression: str = "lossless"
    quality: int = 80
    dpi: int = 200
    ocr_language: str = "kor+eng"
    output_format: str = "pdf"
    force_ocr: bool = False
    psm: int = 3


@dataclass
class Word:
    text: str
    left: float
    top: float
    right: float
    bottom: float


@dataclass
class PageContent:
    text: str
    words: list[Word]
    width: float
    height: float


OPERATIONS = {"merge", "split", "extract", "delete", "rotate", "images_to_pdf",
              "pdf_to_images", "compress", "recognize"}


def check_cancel(cancel):
    if cancel is not None and cancel.is_set():
        raise Cancelled()


def page_indices(expression, count, *, allow_empty=True):
    """1부터 시작하는 범위를 검증하고 사용자가 쓴 페이지 순서를 유지한다."""
    value = expression.strip()
    if count < 1:
        raise PdfToolError("empty_pdf")
    if not value:
        if allow_empty:
            return list(range(count))
        raise PdfToolError("page_range")
    if len(value) > 10000:
        raise PdfToolError("page_range")
    result = []
    for token in value.split(","):
        match = re.fullmatch(r"\s*([1-9][0-9]*)(?:\s*-\s*([1-9][0-9]*))?\s*", token)
        if not match:
            raise PdfToolError("page_range")
        start = int(match.group(1))
        end = int(match.group(2) or start)
        if not (1 <= start <= end <= count):
            raise PdfToolError("page_range")
        for number in range(start - 1, end):
            if number not in result:
                result.append(number)
    return result


def reader_for(path):
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            raise PdfToolError("encrypted_pdf")
        if not reader.pages:
            raise PdfToolError("empty_pdf")
        return reader
    except PdfToolError:
        raise
    except Exception as exc:
        raise PdfToolError("invalid_pdf", str(exc)) from exc


def tesseract_paths():
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    folder = base / "vendor" / "tesseract"
    executable = folder / "tesseract.exe"
    data = folder / "tessdata"
    if not executable.is_file() or not data.is_dir():
        raise PdfToolError("ocr_missing")
    return executable, data


class _ProcessTree:
    """Windows Job Object를 닫아 OCR 프로세스와 그 자식을 함께 종료한다."""
    def __init__(self, process):
        self.handle = None
        if os.name != "nt":
            return

        class Basic(ctypes.Structure):
            _fields_ = [("PerProcessUserTimeLimit", ctypes.c_longlong),
                        ("PerJobUserTimeLimit", ctypes.c_longlong),
                        ("LimitFlags", wintypes.DWORD), ("MinimumWorkingSetSize", ctypes.c_size_t),
                        ("MaximumWorkingSetSize", ctypes.c_size_t), ("ActiveProcessLimit", wintypes.DWORD),
                        ("Affinity", ctypes.c_size_t), ("PriorityClass", wintypes.DWORD),
                        ("SchedulingClass", wintypes.DWORD)]

        class Io(ctypes.Structure):
            _fields_ = [(name, ctypes.c_ulonglong) for name in
                        ("ReadOperationCount", "WriteOperationCount", "OtherOperationCount",
                         "ReadTransferCount", "WriteTransferCount", "OtherTransferCount")]

        class Extended(ctypes.Structure):
            _fields_ = [("BasicLimitInformation", Basic), ("IoInfo", Io),
                        ("ProcessMemoryLimit", ctypes.c_size_t), ("JobMemoryLimit", ctypes.c_size_t),
                        ("PeakProcessMemoryUsed", ctypes.c_size_t), ("PeakJobMemoryUsed", ctypes.c_size_t)]

        self.kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        self.kernel.CreateJobObjectW.argtypes = [ctypes.c_void_p, wintypes.LPCWSTR]
        self.kernel.CreateJobObjectW.restype = wintypes.HANDLE
        self.kernel.SetInformationJobObject.argtypes = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, wintypes.DWORD]
        self.kernel.AssignProcessToJobObject.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
        self.kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        self.handle = self.kernel.CreateJobObjectW(None, None)
        if not self.handle:
            raise ctypes.WinError(ctypes.get_last_error())
        info = Extended()
        info.BasicLimitInformation.LimitFlags = 0x2000
        if not self.kernel.SetInformationJobObject(self.handle, 9, ctypes.byref(info), ctypes.sizeof(info)):
            self.close()
            raise ctypes.WinError(ctypes.get_last_error())
        if not self.kernel.AssignProcessToJobObject(self.handle, int(process._handle)):
            self.close()
            raise ctypes.WinError(ctypes.get_last_error())

    def close(self):
        if self.handle:
            self.kernel.CloseHandle(self.handle)
            self.handle = None


def run_ocr_process(argv, work, cancel, timeout=180):
    check_cancel(cancel)
    log = Path(work) / "ocr-process.log"
    environment = os.environ.copy()
    environment["OMP_THREAD_LIMIT"] = "2"
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    process = None
    tree = None
    with log.open("wb") as stderr:
        try:
            process = subprocess.Popen(list(map(str, argv)), cwd=work, env=environment,
                                       stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=stderr,
                                       shell=False, creationflags=flags, start_new_session=os.name != "nt")
            tree = _ProcessTree(process)
            deadline = time.monotonic() + timeout
            while process.poll() is None:
                check_cancel(cancel)
                if time.monotonic() >= deadline:
                    raise PdfToolError("ocr_timeout")
                if cancel is not None:
                    cancel.wait(0.05)
                else:
                    time.sleep(0.05)
            if process.returncode:
                raise PdfToolError("ocr_failed", f"exit={process.returncode}")
        finally:
            if process is not None and process.poll() is None:
                if tree is not None and tree.handle:
                    tree.close()
                elif os.name != "nt":
                    os.killpg(process.pid, signal.SIGKILL)
                else:
                    process.kill()
                process.wait(timeout=10)
            if tree is not None:
                tree.close()


def parse_tsv(path):
    words = []
    lines = {}
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream, delimiter="\t"):
            text = row.get("text", "").strip()
            if not text or row.get("level") != "5":
                continue
            left, top, width, height = [float(row[key]) for key in ("left", "top", "width", "height")]
            words.append(Word(text, left, top, left + width, top + height))
            key = tuple(row[key] for key in ("page_num", "block_num", "par_num", "line_num"))
            lines.setdefault(key, []).append(text)
    return words, "\n".join(" ".join(line) for line in lines.values())


def image_ocr(image_path, work, request, cancel, *, searchable=False, staged_data=None):
    executable, data = tesseract_paths()
    if request.ocr_language not in {"eng", "kor", "kor+eng"}:
        raise PdfToolError("ocr_language")
    for language in request.ocr_language.split("+"):
        if not (data / (language + ".traineddata")).is_file():
            raise PdfToolError("ocr_missing", language)
    # 배포 C++ 런타임은 한글 절대 데이터 경로에서 실패하므로 파일 인자는 ASCII 상대 경로로 전달한다.
    work = Path(work)
    staged_data = Path(staged_data) if staged_data is not None else work / "tessdata"
    staged_data.mkdir(exist_ok=True)
    for name in [language + ".traineddata" for language in request.ocr_language.split("+")] + ["pdf.ttf"]:
        if not (staged_data / name).is_file():
            shutil.copy2(data / name, staged_data / name)
    local_input = work / "input.png"
    if Path(image_path).resolve() != local_input.resolve():
        shutil.copy2(image_path, local_input)
    output = work / "recognized"
    argv = [executable, "input.png", "recognized", "--tessdata-dir", os.path.relpath(staged_data, work),
            "-l", request.ocr_language, "--oem", "1", "--psm", str(request.psm), "--dpi", str(request.dpi),
            "-c", "tessedit_create_tsv=1", "-c", "tessedit_create_txt=1"]
    if searchable:
        argv += ["-c", "tessedit_create_pdf=1"]
    run_ocr_process(argv, work, cancel)
    tsv = output.with_suffix(".tsv")
    if not tsv.is_file():
        raise PdfToolError("ocr_failed", "TSV output missing")
    words, text = parse_tsv(tsv)
    with Image.open(image_path) as image:
        width, height = image.size
    return PageContent(text, words, width, height), output.with_suffix(".pdf")


@_serialized_pdfium
def pdf_text_pages(path, cancel=None):
    # 글자의 실제 좌표를 읽어 PDF 표를 단순 문자열 한 셀로 축소하지 않는다.
    reader_for(path)
    document = pdfium.PdfDocument(str(path))
    result = []
    try:
        for number in range(len(document)):
            check_cancel(cancel)
            page = document[number]
            textpage = page.get_textpage()
            try:
                text = textpage.get_text_range()
                height = page.get_height()
                words = []
                fragment = ""
                boxes = []

                def flush():
                    nonlocal fragment, boxes
                    if fragment and boxes:
                        words.append(Word(fragment, min(b[0] for b in boxes), min(b[1] for b in boxes),
                                          max(b[2] for b in boxes), max(b[3] for b in boxes)))
                    fragment, boxes = "", []

                for index in range(textpage.count_chars()):
                    char = textpage.get_text_range(index, 1)
                    if not char or char.isspace():
                        flush()
                        continue
                    left, bottom, right, top = textpage.get_charbox(index)
                    box = (left, height - top, right, height - bottom)
                    if boxes and (box[0] - boxes[-1][2] > max(10, box[3] - box[1]) or
                                  abs(box[1] - boxes[-1][1]) > max(4, (box[3] - box[1]) * 0.65)):
                        flush()
                    fragment += char
                    boxes.append(box)
                flush()
                result.append(PageContent(text, words, page.get_width(), height))
            finally:
                textpage.close()
                page.close()
    finally:
        document.close()
    return result


def structured_rows(content):
    """같은 높이의 단어와 큰 가로 간격을 이용해 행·열을 추론한다."""
    words = sorted(content.words, key=lambda word: (word.top, word.left))
    if not words:
        return []
    typical = statistics.median(max(1, word.bottom - word.top) for word in words)
    lines = []
    for word in words:
        center = (word.top + word.bottom) / 2
        line = next((line for line in reversed(lines[-4:]) if abs(line[0] - center) <= typical * 0.6), None)
        if line is None:
            line = [center, []]
            lines.append(line)
        line[1].append(word)
    raw_rows = []
    for _, line in sorted(lines, key=lambda value: value[0]):
        cells = []
        for word in sorted(line, key=lambda value: value.left):
            if cells and word.left - cells[-1][2] <= typical * 1.7:
                cells[-1][0] += " " + word.text
                cells[-1][2] = max(cells[-1][2], word.right)
            else:
                cells.append([word.text, word.left, word.right])
        raw_rows.append(cells)
    # 여러 행에서 반복되는 시작 위치를 열 앵커로 삼는다. 제목/본문 한 줄은 별도로 유지한다.
    anchors = []
    for row in raw_rows:
        if len(row) < 2:
            continue
        for _, left, _ in row:
            if not any(abs(left - anchor) < typical * 1.5 for anchor in anchors):
                anchors.append(left)
    anchors.sort()
    result = []
    for row in raw_rows:
        if len(row) == 1 or len(anchors) < 2:
            result.append([cell[0] for cell in row])
            continue
        values = [""] * len(anchors)
        for text, left, _ in row:
            column = min(range(len(anchors)), key=lambda i: abs(anchors[i] - left))
            values[column] = (values[column] + " " + text).strip()
        result.append(values)
    return result


def export_document(contents, destination, kind):
    if kind == "txt":
        Path(destination).write_text("\n\n\f\n\n".join(page.text for page in contents), encoding="utf-8")
    elif kind == "docx":
        from docx import Document
        document = Document()
        for index, content in enumerate(contents):
            if index:
                document.add_page_break()
            rows = structured_rows(content)
            group = []

            def flush_table():
                if not group:
                    return
                width = max(map(len, group))
                table = document.add_table(rows=len(group), cols=width)
                table.style = "Table Grid"
                for row_index, row in enumerate(group):
                    for column, value in enumerate(row):
                        table.cell(row_index, column).text = value
                group.clear()

            for row in rows:
                if len(row) > 1:
                    group.append(row)
                else:
                    flush_table()
                    document.add_paragraph(row[0] if row else "")
            flush_table()
        document.save(destination)
    elif kind == "xlsx":
        import xlsxwriter
        workbook = xlsxwriter.Workbook(str(destination), {"strings_to_formulas": False, "strings_to_urls": False})
        try:
            style = workbook.add_format({"text_wrap": True, "valign": "top"})
            for index, content in enumerate(contents):
                sheet = workbook.add_worksheet(f"Page {index + 1}")
                sheet.set_column(0, 30, 22)
                for row_index, row in enumerate(structured_rows(content)):
                    for column, text in enumerate(row):
                        sheet.write_string(row_index, column, text, style)
        finally:
            workbook.close()
    else:
        raise PdfToolError("output_format")


@_serialized_pdfium
def render_page(document, index, destination, dpi, cancel=None):
    check_cancel(cancel)
    page = document[index]
    bitmap = None
    try:
        scale = dpi / 72
        if page.get_width() * page.get_height() * scale * scale > 100_000_000:
            raise PdfToolError("page_too_large")
        bitmap = page.render(scale=scale)
        image = bitmap.to_pil().copy().convert("RGB")
        image.save(destination, dpi=(dpi, dpi))
    finally:
        if bitmap is not None:
            bitmap.close()
        page.close()


def publish(produced, source, label, cancel=None):
    check_cancel(cancel)
    source = Path(source)
    produced = Path(produced)
    for number in range(1, 100000):
        suffix = "" if number == 1 else f" {number}"
        destination = source.parent / f"{source.stem}-{label}{suffix}{produced.suffix}"
        try:
            # 배타적 생성으로 기존 원본·이전 결과를 덮어쓰지 않는다.
            stream = destination.open("xb")
        except FileExistsError:
            continue
        try:
            with stream, produced.open("rb") as original:
                while block := original.read(1024 * 1024):
                    check_cancel(cancel)
                    stream.write(block)
            check_cancel(cancel)
            return str(destination)
        except BaseException:
            destination.unlink(missing_ok=True)
            raise
    raise PdfToolError("output_name")


def _write_pdf(writer, path):
    if not writer.pages:
        raise PdfToolError("all_pages_removed")
    with Path(path).open("wb") as stream:
        writer.write(stream)
    writer.close()


@_serialized_pdfium
def _recognize(source, request, work, cancel, progress):
    source = Path(source)
    is_pdf = source.suffix.lower() == ".pdf"
    text_pages = pdf_text_pages(source, cancel) if is_pdf else []
    searchable = request.output_format == "pdf"
    contents, pdf_pages = [], []
    document = pdfium.PdfDocument(str(source)) if is_pdf else None
    image_input = Image.open(source) if not is_pdf else None
    count = len(document) if document is not None else getattr(image_input, "n_frames", 1)
    chosen = page_indices(request.pages, count)
    try:
        for number, index in enumerate(chosen):
            check_cancel(cancel)
            progress(number, len(chosen), f"{source.name} · {index + 1}")
            if is_pdf and not searchable and not request.force_ocr and text_pages[index].text.strip():
                contents.append(text_pages[index])
                continue
            page_work = Path(work) / f"page-{index + 1}"
            page_work.mkdir()
            raster = page_work / "input.png"
            if document is not None:
                render_page(document, index, raster, request.dpi, cancel)
            else:
                image_input.seek(index)
                image = ImageOps.exif_transpose(image_input).convert("RGB")
                image.save(raster, dpi=(request.dpi, request.dpi))
            content, pdf = image_ocr(raster, page_work, request, cancel, searchable=searchable,
                                     staged_data=Path(work) / "tessdata")
            contents.append(content)
            if searchable:
                if not pdf.is_file():
                    raise PdfToolError("ocr_failed", "PDF output missing")
                pdf_pages.append(pdf)
        output = Path(work) / ("recognized." + request.output_format)
        if searchable:
            writer = PdfWriter()
            for pdf in pdf_pages:
                check_cancel(cancel)
                writer.append(str(pdf))
            _write_pdf(writer, output)
        else:
            export_document(contents, output, request.output_format)
        return output
    finally:
        if document is not None:
            document.close()
        if image_input is not None:
            image_input.close()


@_serialized_pdfium
def execute(request, cancel=None, progress=None, result=None):
    """결과별 (입력, 출력, 오류) 콜백. 취소/실패 시 완성한 앞선 결과는 유지한다."""
    cancel = cancel if cancel is not None else threading.Event()
    progress = progress or (lambda done, total, label: None)
    result = result or (lambda source, output, error: None)
    if request.operation not in OPERATIONS:
        raise PdfToolError("operation")
    paths = []
    for raw in request.paths:
        with source_context(str(Path(raw).resolve())):
            path = Path(raw).resolve(strict=True)
            if not path.is_file():
                raise PdfToolError("input_files")
            paths.append(path)
    if not paths:
        raise PdfToolError("input_files")
    if not 60 <= request.dpi <= 600 or not 20 <= request.quality <= 95:
        raise PdfToolError("quality")
    if request.rotation not in (90, 180, 270):
        raise PdfToolError("rotation")
    if request.psm not in (3, 6, 11):
        raise PdfToolError("ocr_layout")
    if request.output_format not in ("pdf", "txt", "docx", "xlsx", "png", "jpg"):
        raise PdfToolError("output_format")
    outputs = []

    def complete(produced, source, label):
        destination = publish(produced, source, label, cancel)
        outputs.append(destination)
        result(str(source), destination, "")

    check_cancel(cancel)
    if request.operation in ("merge", "images_to_pdf"):
        source = paths[0]
        with tempfile.TemporaryDirectory(prefix=".flipring-pdf-", dir=source.parent) as folder:
            work = Path(folder)
            output = work / "combined.pdf"
            if request.operation == "merge":
                writer = PdfWriter()
                for number, path in enumerate(paths):
                    check_cancel(cancel)
                    progress(number, len(paths), path.name)
                    with source_context(path):
                        reader = reader_for(path)
                    writer.append(reader)
                _write_pdf(writer, output)
            else:
                from reportlab.pdfgen import canvas
                from reportlab.lib.utils import ImageReader
                document = canvas.Canvas(str(output))
                for number, path in enumerate(paths):
                    progress(number, len(paths), path.name)
                    with source_context(path), Image.open(path) as source_image:
                        for frame in ImageSequence.Iterator(source_image):
                            check_cancel(cancel)
                            image = ImageOps.exif_transpose(frame).convert("RGB")
                            width, height = (dimension * 72 / request.dpi for dimension in image.size)
                            document.setPageSize((width, height))
                            document.drawImage(ImageReader(image), 0, 0, width=width, height=height)
                            document.showPage()
                document.save()
            complete(output, source, request.operation)
        progress(len(paths), len(paths), "")
        return outputs

    for number, source in enumerate(paths):
        check_cancel(cancel)
        progress(number, len(paths), source.name)
        with source_context(source), tempfile.TemporaryDirectory(prefix=".flipring-pdf-", dir=source.parent) as folder:
            work = Path(folder)
            if request.operation == "recognize":
                output = _recognize(source, request, work, cancel, progress)
                complete(output, source, "ocr" if request.force_ocr or source.suffix.lower() != ".pdf" else "recognized")
                continue
            reader = reader_for(source)
            selected = page_indices(request.pages, len(reader.pages), allow_empty=request.operation != "delete")
            if request.operation == "pdf_to_images":
                document = pdfium.PdfDocument(str(source))
                try:
                    for index in selected:
                        raster = work / f"page-{index + 1}.png"
                        render_page(document, index, raster, request.dpi, cancel)
                        if request.output_format == "jpg":
                            with Image.open(raster) as image:
                                raster = raster.with_suffix(".jpg")
                                image.convert("RGB").save(raster, quality=request.quality)
                        complete(raster, source, f"page-{index + 1}")
                finally:
                    document.close()
                continue
            if request.operation == "split":
                for index in selected:
                    check_cancel(cancel)
                    writer = PdfWriter()
                    writer.add_page(reader.pages[index])
                    output = work / f"page-{index + 1}.pdf"
                    _write_pdf(writer, output)
                    complete(output, source, f"page-{index + 1}")
                continue
            writer = PdfWriter()
            indices = selected if request.operation == "extract" else list(range(len(reader.pages)))
            if request.operation == "delete":
                indices = [index for index in indices if index not in selected]
            for index in indices:
                check_cancel(cancel)
                page = writer.add_page(reader.pages[index])
                if request.operation == "rotate" and index in selected:
                    page.rotate(request.rotation)
                if request.operation == "compress":
                    if request.compression not in ("lossless", "images"):
                        raise PdfToolError("compression")
                    if request.compression == "images":
                        cap = max(1, int(max(float(page.mediabox.width), float(page.mediabox.height)) / 72 * request.dpi))
                        for image in list(page.images):
                            check_cancel(cancel)
                            resized = image.image.copy()
                            resized.thumbnail((cap, cap))
                            image.replace(resized, quality=request.quality)
                    page.compress_content_streams(level=6)
            if request.operation == "compress":
                writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
            output = work / "result.pdf"
            _write_pdf(writer, output)
            complete(output, source, request.operation)
    progress(len(paths), len(paths), "")
    return outputs
