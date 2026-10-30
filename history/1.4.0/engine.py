"""PUVKID의 Windows 파일 변환 엔진. 원본은 읽기만 하고 결과를 새 이름으로 저장한다."""
from __future__ import annotations

import csv
import html
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tarfile
import tempfile
import zipfile

from PIL import Image, ImageOps
import pi_heif
import pypdfium2 as pdfium
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

pi_heif.register_heif_opener()
TABLE = {
    "이미지": ("jpg png webp heic tiff avif bmp gif svg", "jpg png webp tiff avif bmp pdf"),
    "오디오": ("mp3 m4a wav flac ogg opus aiff wma", "mp3 m4a wav flac ogg opus aiff wma"),
    "비디오": ("mp4 mov mkv webm avi wmv", "mp4 mov mkv webm avi wmv gif mp3"),
    "PDF": ("pdf", "docx jpg png epub txt md"),
    "문서": ("docx doc odt rtf md txt epub html", "docx md txt epub html pdf jpg png"),
    "프레젠테이션": ("pptx ppt odp", "pdf jpg png"),
    "스프레드시트": ("xlsx xls ods csv tsv", "xlsx ods csv tsv pdf"),
    "자막": ("srt vtt", "srt vtt txt"),
    "압축": ("zip tar tar.gz rar", "zip tar tar.gz"),
    "폴더": ("folder", "zip tar tar.gz"),
}
ALIASES = dict(jpeg="jpg", jpe="jpg", tif="tiff", heif="heic", htm="html", markdown="md", aif="aiff", tgz="tar.gz")
TOOLS = {
    "ffmpeg": ("FFmpeg", "Gyan.FFmpeg", "오디오 · 비디오"),
    "pandoc": ("Pandoc", "JohnMacFarlane.Pandoc", "전자책 · 문서"),
    "soffice": ("LibreOffice", "TheDocumentFoundation.LibreOffice", "Office 문서 · 발표 · 표"),
}


def asset(name):
    return Path(__file__).resolve().parent / "assets" / name


def classify(path):
    p = Path(path)
    ext = "folder" if p.is_dir() else "tar.gz" if p.name.lower().endswith((".tar.gz", ".tgz")) else p.suffix.lower().lstrip(".")
    ext = ALIASES.get(ext, ext)
    for kind, (sources, _) in TABLE.items():
        if ext in sources.split():
            return kind, ext
    raise ValueError(f"지원하지 않는 파일입니다: {p.name}")


def stem(path):
    p = Path(path)
    if p.is_dir():
        return p.name
    return p.name[:-7] if p.name.lower().endswith(".tar.gz") else p.stem


def targets(paths):
    common = None
    for p in paths:
        try:
            kind, ext = classify(p)
        except ValueError:
            return []
        options = [x for x in TABLE[kind][1].split() if x != ext]
        common = options if common is None else [x for x in common if x in options]
    return common or []


def tool(name):
    # 현재 프로세스를 시작한 뒤 설치한 도구도 다음 드래그에서 탐색한다.
    exe = name + ".exe"
    p = shutil.which(exe)
    if p:
        return p
    local = Path(os.environ.get("LOCALAPPDATA", ""))
    roots = [local / "Microsoft/WinGet/Links", Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "LibreOffice/program", local / "Pandoc", Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "Pandoc"]
    for root in roots:
        if (root / exe).is_file():
            return str(root / exe)
    packages = local / "Microsoft/WinGet/Packages"
    for prefix in {"ffmpeg": ["Gyan.FFmpeg*"], "pandoc": ["JohnMacFarlane.Pandoc*"], "soffice": []}.get(name, []):
        for root in packages.glob(prefix):
            found = next(root.rglob(exe), None)
            if found:
                return str(found)
    return None


def requirement(path, dst):
    kind, src = classify(path)
    needed = None
    if kind in ("오디오", "비디오"):
        needed = "ffmpeg"
    elif kind in ("프레젠테이션", "스프레드시트"):
        if not (src in ("csv", "tsv") and dst in ("csv", "tsv")):
            needed = "soffice"
    elif kind == "문서":
        if src in ("doc", "odt", "rtf"):
            if not tool("soffice"): return "LibreOffice"
            if dst in ("md", "epub"): needed = "pandoc"
        elif src == "epub" or dst == "epub" or (src == "html" and dst != "txt") or (src == "md" and dst != "txt") or (src == "docx" and dst in ("html", "md")):
            if dst in ("pdf", "jpg", "png") and not tool("soffice"):
                return "LibreOffice"
            needed = "pandoc"
        elif src == "docx" and dst in ("pdf", "jpg", "png"):
            needed = "soffice"
    elif kind == "PDF" and dst == "epub":
        needed = "pandoc"
    return TOOLS[needed][0] if needed and not tool(needed) else None


def run(exe, args, cwd=None):
    if not exe:
        raise RuntimeError("변환 도구를 찾을 수 없습니다. 설정의 도우미 도구를 확인해 주세요.")
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    # 외부 도구가 앱에 묶인 Qt/Python DLL을 잘못 로드하지 않도록 생성 시 검색 경로를 복구한다.
    frozen = os.name == "nt" and getattr(sys, "frozen", False)
    if frozen:
        import ctypes
        ctypes.windll.kernel32.SetDllDirectoryW(None)
    try:
        process = subprocess.Popen([str(exe), *map(str, args)], cwd=cwd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE, creationflags=flags)
    finally:
        if frozen: ctypes.windll.kernel32.SetDllDirectoryW(sys._MEIPASS)
    try: stdout, stderr = process.communicate(timeout=1800)
    except subprocess.TimeoutExpired:
        process.kill(); process.communicate(); raise RuntimeError("변환 제한 시간 30분을 초과했습니다.")
    if process.returncode:
        detail = (stderr or stdout).decode("utf-8", "replace")[-1200:]
        raise RuntimeError(f"{Path(exe).stem} 변환 실패 ({process.returncode})\n{detail}")
    return stdout


def text(path):
    data = Path(path).read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp949", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    raise ValueError("텍스트 인코딩을 확인해 주세요.")


def write_text_pdf(content, output):
    font = "PUVKIDText"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, str(asset("NanumGothicBold.ttf"))))
    c = canvas.Canvas(str(output), pagesize=(612, 792))
    c.setTitle(output.stem)
    c.setFont(font, 11)
    y = 738
    # 사용자 파일의 글자는 자르지 않고 폭에 맞춰 줄바꿈한다.
    for line in content.expandtabs(4).splitlines() or [""]:
        fragment = ""
        for char in line:
            if fragment and pdfmetrics.stringWidth(fragment + char, font, 11) > 504:
                c.drawString(54, y, fragment)
                y -= 17
                fragment = ""
                if y < 54:
                    c.showPage(); c.setFont(font, 11); y = 738
            fragment += char
        c.drawString(54, y, fragment)
        y -= 17
        if y < 54:
            c.showPage(); c.setFont(font, 11); y = 738
    c.save()


def render_pdf(source, dst, work, base):
    doc = pdfium.PdfDocument(str(source))
    try:
        if not len(doc):
            raise ValueError("페이지가 없는 PDF입니다.")
        out = work / f"{base} ({dst})" if len(doc) > 1 else work / f"{base}.{dst}"
        if len(doc) > 1:
            out.mkdir()
        for i in range(len(doc)):
            page = doc[i]
            bitmap = page.render(scale=2)
            try:
                image = bitmap.to_pil().copy().convert("RGB")
                dest = out / f"{base}-{i + 1:0{len(str(len(doc)))}d}.{dst}" if len(doc) > 1 else out
                image.save(dest, quality=90)
            finally:
                bitmap.close(); page.close()
        return out
    finally:
        doc.close()


def office(source, dst, work):
    folder = work / ("office-" + next(tempfile._get_candidate_names()))
    folder.mkdir()
    profile = (folder / "profile").as_uri()
    filters = {"csv": "csv:Text - txt - csv (StarCalc):44,34,76", "tsv": "csv:Text - txt - csv (StarCalc):9,34,76", "txt": "txt:Text (encoded):UTF8", "xlsx": "xlsx:Calc MS Excel 2007 XML"}
    extra = ["--infilter=CSV:" + ("9" if source.suffix.lower() == ".tsv" else "44") + ",34,76"] if source.suffix.lower() in (".csv", ".tsv") else []
    run(tool("soffice"), [f"-env:UserInstallation={profile}", "--headless", "--norestore", *extra, "--convert-to", filters.get(dst, dst), "--outdir", folder, source])
    result = folder / f"{source.stem}.{'csv' if dst == 'tsv' else dst}"
    if not result.is_file():
        raise RuntimeError(f"LibreOffice에서 {dst.upper()} 결과를 만들지 못했습니다.")
    if dst == "tsv":
        result = result.rename(result.with_suffix(".tsv"))
    return result


def pandoc(source, src, dst, output):
    reader = {"md": "markdown", "txt": "markdown"}.get(src, src)
    writer = {"md": "gfm", "txt": "plain"}.get(dst, dst)
    run(tool("pandoc"), [source, "--from", reader, "--to", writer, "--standalone", "--wrap=none", "--metadata", f"title={source.stem}", "-o", output], source.parent)
    return output


def document(source, src, dst, work, base):
    output = work / f"{base}.{dst}"
    if src in ("doc", "odt", "rtf"):
        if dst in ("pdf", "docx", "txt", "html"):
            return office(source, dst, work)
        if dst in ("jpg", "png"):
            return render_pdf(office(source, "pdf", work), dst, work, base)
        source = office(source, "docx", work); src = "docx"
    if src == "docx" and dst in ("pdf", "jpg", "png"):
        pdf = office(source, "pdf", work)
        return pdf if dst == "pdf" else render_pdf(pdf, dst, work, base)
    if src == "epub" or dst == "epub" or (src == "html" and dst != "txt") or (src == "md" and dst != "txt") or (src == "docx" and dst in ("md", "html")):
        if dst in ("pdf", "jpg", "png"):
            intermediate = work / "pandoc-source.docx"
            pandoc(source, src, "docx", intermediate)
            pdf = office(intermediate, "pdf", work)
            return pdf if dst == "pdf" else render_pdf(pdf, dst, work, base)
        return pandoc(source, src, dst, output)
    if src == "docx":
        doc = Document(source)
        content = "\n".join([p.text for p in doc.paragraphs] + ["\t".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    else:
        content = text(source)
        if src == "html":
            from html.parser import HTMLParser
            class Extract(HTMLParser):
                def __init__(self): super().__init__(); self.parts = []; self.skip = 0
                def handle_starttag(self, tag, attrs):
                    if tag in ("script", "style"): self.skip += 1
                    if tag in ("p", "br", "div", "li", "h1", "h2"): self.parts.append("\n")
                def handle_endtag(self, tag):
                    if tag in ("script", "style"): self.skip = max(0, self.skip - 1)
                def handle_data(self, data):
                    if not self.skip: self.parts.append(data)
            parser = Extract(); parser.feed(content); content = "".join(parser.parts)
    if dst in ("txt", "md"):
        output.write_text(content, encoding="utf-8")
    elif dst == "html":
        output.write_text('<!doctype html><meta charset="utf-8"><pre>' + html.escape(content) + "</pre>", encoding="utf-8")
    elif dst == "docx":
        d = Document()
        for line in content.splitlines(): d.add_paragraph(line)
        d.save(output)
    elif dst in ("pdf", "jpg", "png"):
        pdf = work / (base + ".pdf")
        write_text_pdf(content, pdf)
        return pdf if dst == "pdf" else render_pdf(pdf, dst, work, base)
    else:
        raise ValueError("지원하지 않는 문서 출력입니다.")
    return output


def subtitles(source, dst, output):
    raw = text(source).replace("\r\n", "\n").replace("\r", "\n")
    cues = []
    for block in re.split(r"\n\s*\n", raw.strip()):
        lines = block.splitlines()
        for i, line in enumerate(lines):
            m = re.match(r"\s*((?:\d+:)?\d{2}:\d{2}[.,]\d{3})\s*-->\s*((?:\d+:)?\d{2}:\d{2}[.,]\d{3})", line)
            if m:
                times = [v.replace(",", ".") if v.count(":") == 2 else "00:" + v.replace(",", ".") for v in m.groups()]
                cues.append((*times, lines[i + 1:])); break
    if not cues:
        raise ValueError("자막 시간 정보를 찾지 못했습니다.")
    if dst == "txt":
        clean = []
        for _, _, lines in cues:
            for line in lines:
                value = html.unescape(re.sub(r"<[^>]*>|\{\\[^}]*\}", "", line)).strip()
                if value and (not clean or clean[-1] != value): clean.append(value)
        value = "\n".join(clean) + "\n"
    else:
        blocks = []
        for i, (start, end, lines) in enumerate(cues, 1):
            if dst == "srt": start, end = start.replace(".", ","), end.replace(".", ",")
            blocks.append((f"{i}\n" if dst == "srt" else "") + f"{start} --> {end}\n" + "\n".join(lines))
        value = ("WEBVTT\n\n" if dst == "vtt" else "") + "\n\n".join(blocks) + "\n"
    output.write_text(value, encoding="utf-8")
    return output


def safe_member(root, name):
    # 압축 안의 절대 경로·상위 경로·Windows 장치 경로를 거부한다.
    clean = name.replace("\\", "/")
    if ":" in clean or clean.startswith("/") or ".." in clean.split("/"):
        raise ValueError("압축 파일에 안전하지 않은 경로가 있습니다.")
    destination = (root / clean).resolve()
    if not destination.is_relative_to(root.resolve()):
        raise ValueError("압축 해제 경로가 작업 폴더를 벗어납니다.")
    return destination


def archive(source, src, dst, output, work):
    import stat
    if src == "folder":
        root, prefix = source, source.name
    else:
        root = work / "extracted"; root.mkdir(); prefix = ""
        if src == "zip":
            with zipfile.ZipFile(source) as z:
                for entry in z.infolist():
                    safe_member(root, entry.filename)
                    if stat.S_ISLNK(entry.external_attr >> 16): raise ValueError("심볼릭 링크 압축은 지원하지 않습니다.")
                z.extractall(root)
        elif src == "rar":
            import rarfile
            sevenzip = Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "7-Zip/7z.exe"
            if sevenzip.is_file(): rarfile.SEVENZIP_TOOL = str(sevenzip)
            bsdtar = Path(os.environ.get("WINDIR", "C:/Windows")) / "System32/tar.exe"
            if bsdtar.is_file(): rarfile.BSDTAR_TOOL = str(bsdtar)
            with rarfile.RarFile(source) as z:
                for entry in z.infolist():
                    safe_member(root, entry.filename)
                    if entry.is_symlink(): raise ValueError("심볼릭 링크 압축은 지원하지 않습니다.")
                try: z.extractall(root)
                except rarfile.RarCannotExec as exc: raise RuntimeError("RAR 해제에 UnRAR 또는 7-Zip 도구가 필요합니다.") from exc
        else:
            with tarfile.open(source) as z:
                for entry in z.getmembers():
                    safe_member(root, entry.name)
                    if not (entry.isfile() or entry.isdir()): raise ValueError("링크·특수 파일 압축은 지원하지 않습니다.")
                z.extractall(root, filter="data")
    entries = list(root.rglob("*"))
    if any(p.is_symlink() or (getattr(p.stat(follow_symlinks=False), "st_file_attributes", 0) & 0x400) for p in entries):
        raise ValueError("링크·접합점이 포함된 폴더는 압축하지 않습니다.")
    if dst == "zip":
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as z:
            if prefix: z.writestr(prefix + "/", "")
            for p in entries: z.write(p, str(Path(prefix) / p.relative_to(root)))
    else:
        with tarfile.open(output, "w:gz" if dst == "tar.gz" else "w") as z:
            if prefix: z.add(root, arcname=prefix, recursive=False)
            for p in entries: z.add(p, arcname=str(Path(prefix) / p.relative_to(root)), recursive=False)
    return output


def produce(source, dst, work):
    kind, src = classify(source)
    base = stem(source)
    out = work / f"{base}.{dst}"
    if kind == "이미지":
        if src == "svg":
            from svglib.svglib import svg2rlg
            from reportlab.graphics import renderPDF
            pdf = work / "svg.pdf"
            drawing = svg2rlg(str(source))
            if drawing is None: raise ValueError("SVG를 읽을 수 없습니다.")
            renderPDF.drawToFile(drawing, str(pdf))
            if dst == "pdf": return pdf
            source = render_pdf(pdf, "png", work, "svg-raster")
        with Image.open(source) as original:
            im = ImageOps.exif_transpose(original).copy()
            if dst in ("jpg", "bmp", "pdf"):
                rgba = im.convert("RGBA"); flat = Image.new("RGB", im.size, "white"); flat.paste(rgba, mask=rgba.getchannel("A")); im = flat
            im.save(out, format={"jpg": "JPEG", "tiff": "TIFF"}.get(dst, dst.upper()), quality=90)
        return out
    if kind in ("오디오", "비디오"):
        args = ["-nostdin", "-hide_banner", "-loglevel", "error", "-i", source]
        audio = {"mp3": ["libmp3lame", "-q:a", "2"], "m4a": ["aac", "-b:a", "256k"], "wav": ["pcm_s16le"], "flac": ["flac"], "ogg": ["libvorbis", "-q:a", "6"], "opus": ["libopus", "-b:a", "160k"], "aiff": ["pcm_s16be"], "wma": ["wmav2", "-b:a", "192k"]}
        if kind == "오디오" or dst == "mp3": args += ["-vn", "-c:a", *audio[dst]]
        elif dst in ("mp4", "mov", "mkv"): args += ["-vf", "scale=trunc(iw/2)*2:trunc(ih/2)*2", "-c:v", "libx264", "-crf", "20", "-pix_fmt", "yuv420p", "-c:a", "aac"] + (["-movflags", "+faststart"] if dst != "mkv" else [])
        elif dst == "webm": args += ["-c:v", "libvpx-vp9", "-crf", "32", "-b:v", "0", "-c:a", "libopus"]
        elif dst == "avi": args += ["-c:v", "mpeg4", "-q:v", "3", "-c:a", "libmp3lame"]
        elif dst == "wmv": args += ["-c:v", "wmv2", "-b:v", "4M", "-c:a", "wmav2"]
        elif dst == "gif": args += ["-an", "-filter_complex", "fps=12,scale=640:-2:force_original_aspect_ratio=decrease,split[a][b];[a]palettegen[p];[b][p]paletteuse", "-loop", "0"]
        run(tool("ffmpeg"), [*args, out]); return out
    if kind == "PDF":
        if dst in ("jpg", "png"): return render_pdf(source, dst, work, base)
        doc = pdfium.PdfDocument(str(source))
        try:
            pages = []
            for i in range(len(doc)):
                page = doc[i]; tp = page.get_textpage()
                try: pages.append(tp.get_text_range())
                finally: tp.close(); page.close()
        finally: doc.close()
        content = ("\n\n---\n\n" if dst == "md" else "\n\n").join(pages)
        if dst in ("txt", "md"): out.write_text(content, encoding="utf-8"); return out
        if dst == "docx":
            d = Document()
            for line in content.splitlines(): d.add_paragraph(line)
            d.save(out); return out
        txt = work / f"{base}.txt"; txt.write_text(content, encoding="utf-8"); return pandoc(txt, "txt", dst, out)
    if kind == "문서": return document(source, src, dst, work, base)
    if kind == "프레젠테이션":
        pdf = office(source, "pdf", work)
        return pdf if dst == "pdf" else render_pdf(pdf, dst, work, base)
    if kind == "스프레드시트":
        if src in ("csv", "tsv") and dst in ("csv", "tsv"):
            rows = list(csv.reader(io.StringIO(text(source), newline=""), delimiter="\t" if src == "tsv" else ","))
            with out.open("w", encoding="utf-8-sig", newline="") as f: csv.writer(f, delimiter="\t" if dst == "tsv" else ",").writerows(rows)
            return out
        return office(source, dst, work)
    if kind == "자막": return subtitles(source, dst, out)
    if kind in ("압축", "폴더"): return archive(source, src, dst, out, work)
    raise ValueError("지원하지 않는 변환입니다.")


def convert(path, dst):
    source = Path(path).resolve(strict=True)
    if dst not in targets([source]): raise ValueError("이 파일에 사용할 수 없는 출력 형식입니다.")
    missing = requirement(source, dst)
    if missing: raise RuntimeError(f"{missing} 도구를 설치해 주세요. 설정에서 설치 명령을 복사할 수 있습니다.")
    # 대상과 같은 드라이브에서 준비하고 최종 이동은 덮어쓰기 없이 처리한다.
    with tempfile.TemporaryDirectory(prefix=".puvkid-dragwheel-", dir=source.parent) as folder:
        work = Path(folder)
        produced = produce(source, dst, work)
        if not produced.exists() or (produced.is_file() and produced.stat().st_size == 0 and dst not in ("txt", "md")):
            raise RuntimeError("변환기가 유효한 출력 파일을 만들지 못했습니다.")
        base = stem(source)
        for i in range(1, 100000):
            suffix = "" if i == 1 else f" {i}"
            candidate = source.parent / (base + suffix + (f" ({dst})" if produced.is_dir() else "." + dst))
            if candidate.exists(): continue
            try:
                # Windows의 rename은 대상이 존재하면 실패하므로 원본과 기존 결과를 보존한다.
                produced.rename(candidate)
                return str(candidate)
            except FileExistsError:
                continue
        raise RuntimeError("새 파일 이름을 만들지 못했습니다.")
