"""실제 Windows 변환과 원본 보존을 검증한다. 생성한 시험 파일만 사용한다."""
import csv
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import time
import zipfile
from PIL import Image
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
import engine


def main():
    root = Path(sys.argv[1])
    root.mkdir(parents=True, exist_ok=False)
    files = []
    image = Image.new("RGBA", (48, 36), (180, 80, 30, 120))
    for ext in "jpg png webp tiff avif bmp gif heic".split():
        p = root / f"사진.{ext}"
        if ext == "heic":
            p.write_bytes((Path(__file__).parent / "test-fixtures/sample.heic").read_bytes())
        else: image.convert("RGB").save(p, format={"jpg": "JPEG"}.get(ext, ext.upper()))
        files.append(p)
    svg = root / "벡터.svg"; svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="48" height="36"><rect width="48" height="36" fill="red"/></svg>', encoding="utf-8"); files.append(svg)
    txt = root / "문서.txt"; txt.write_text("PUVKID 변환 시험\n둘째 줄 한글\n", encoding="utf-8"); files.append(txt)
    md = root / "마크다운.md"; md.write_text("# PUVKID\n\n**한글** 본문입니다.", encoding="utf-8"); files.append(md)
    html = root / "웹문서.html"; html.write_text('<!doctype html><meta charset="utf-8"><h1>PUVKID</h1><p>한글 본문</p>', encoding="utf-8"); files.append(html)
    docx = root / "워드.docx"; d = Document(); d.add_paragraph("PUVKID 한글 문서 시험"); d.save(docx); files.append(docx)
    for ext in ("doc", "odt", "rtf"):
        with tempfile.TemporaryDirectory() as work:
            p = engine.office(docx, ext, Path(work)); dest = root / f"레거시.{ext}"; dest.write_bytes(p.read_bytes()); files.append(dest)
    epub = root / "전자책.epub"; engine.pandoc(md, "md", "epub", epub); files.append(epub)
    pdf = root / "한페이지.pdf"; engine.write_text_pdf("PUVKID 한글 PDF 시험", pdf); files.append(pdf)
    multi = root / "여러페이지.pdf"; engine.write_text_pdf("여러 페이지 시험\n" * 100, multi)
    pptx = root / "발표.pptx"; prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = "PUVKID 시험"; prs.save(pptx); files.append(pptx)
    for ext in ("ppt", "odp"):
        with tempfile.TemporaryDirectory() as work:
            p = engine.office(pptx, ext, Path(work)); dest = root / f"발표원본.{ext}"; dest.write_bytes(p.read_bytes()); files.append(dest)
    xlsx = root / "표.xlsx"; wb = Workbook(); wb.active.append(["한글", "값"]); wb.active.append(["PUVKID", 12]); wb.save(xlsx); files.append(xlsx)
    for ext in ("xls", "ods"):
        with tempfile.TemporaryDirectory() as work:
            p = engine.office(xlsx, ext, Path(work)); dest = root / f"표원본.{ext}"; dest.write_bytes(p.read_bytes()); files.append(dest)
    for ext, delim in [("csv", ","), ("tsv", "\t")]:
        p = root / f"구분자.{ext}"
        with p.open("w", encoding="utf-8-sig", newline="") as f: csv.writer(f, delimiter=delim).writerows([["한글", "따옴표\"", "여러\n줄"], ["가", "1", "2"]])
        files.append(p)
    srt = root / "자막.srt"; srt.write_text("1\n00:00:00,000 --> 00:00:01,000\n<b>한글 자막</b>\n\n2\n00:00:01,000 --> 00:00:02,000\n둘째 줄\n", encoding="utf-8"); files.append(srt)
    vtt = root / "웹자막.vtt"; vtt.write_text("WEBVTT\n\n00:00.000 --> 00:01.000\n한글 자막\n", encoding="utf-8"); files.append(vtt)
    folder = root / "압축 폴더"; folder.mkdir(); (folder / "한글.txt").write_text("원본 보존", encoding="utf-8"); (folder / "빈폴더").mkdir(); files.append(folder)
    for ext in ("zip", "tar", "tar.gz"):
        output = Path(engine.convert(folder, ext)); files.append(output)
    ffmpeg = engine.tool("ffmpeg")
    wav = root / "소리.wav"; engine.run(ffmpeg, ["-f", "lavfi", "-i", "sine=frequency=440:duration=0.3", wav]); files.append(wav)
    for ext in "mp3 m4a flac ogg opus aiff wma".split(): files.append(Path(engine.convert(wav, ext)))
    video = root / "영상.mp4"; engine.run(ffmpeg, ["-f", "lavfi", "-i", "color=c=blue:s=64x48:d=0.3", "-f", "lavfi", "-i", "sine=frequency=440:duration=0.3", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", video]); files.append(video)
    for ext in "mov mkv webm avi wmv".split(): files.append(Path(engine.convert(video, ext)))
    results = []
    for source in files:
        before = hashlib.sha256(source.read_bytes()).hexdigest() if source.is_file() else None
        for target in engine.targets([source]):
            start = time.monotonic()
            try:
                output = Path(engine.convert(source, target))
                assert output.exists() and output.parent == source.parent
                if output.is_file() and target in ("jpg", "png", "webp", "tiff", "avif", "bmp"):
                    with Image.open(output) as im: im.verify()
                if output.is_file() and target == "pdf":
                    doc = engine.pdfium.PdfDocument(str(output)); assert len(doc); doc.close()
                if before: assert before == hashlib.sha256(source.read_bytes()).hexdigest()
                results.append({"source": source.name, "target": target, "ok": True, "seconds": round(time.monotonic()-start, 2)})
            except Exception as exc:
                results.append({"source": source.name, "target": target, "ok": False, "error": str(exc)})
            print(json.dumps(results[-1], ensure_ascii=False), flush=True)
        (root.parent / "진행결과.tmp").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    pages = Path(engine.convert(multi, "png")); assert pages.is_dir() and len(list(pages.glob("*.png"))) > 1
    assert engine.targets([srt, txt]) == []
    first = engine.convert(srt, "vtt"); second = engine.convert(srt, "vtt"); assert first != second
    hostile = root / "경로검증.zip"
    with zipfile.ZipFile(hostile, "w") as z: z.writestr("../outside.txt", "거부")
    try: engine.convert(hostile, "tar")
    except ValueError: pass
    else: raise AssertionError("상위 경로를 거부하지 않았습니다")
    summary = {"total": len(results), "passed": sum(r["ok"] for r in results), "failed": [r for r in results if not r["ok"]], "results": results}
    (root.parent / "변환시험.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k:v for k,v in summary.items() if k != "results"}, ensure_ascii=False))
    return bool(summary["failed"])


if __name__ == "__main__": sys.exit(main())
