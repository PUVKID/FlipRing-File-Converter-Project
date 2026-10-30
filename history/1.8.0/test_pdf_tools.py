"""생성한 PDF·이미지로 원본 보존, 페이지 편집, 표와 실제 한영 OCR을 검증한다."""
from __future__ import annotations

import argparse
import ctypes
import hashlib
import json
import os
from pathlib import Path
import sys
import threading
import time
import unittest
from unittest.mock import patch

from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from openpyxl import load_workbook
import pdf_tools as tools


ROOT = None
OCR = False
GUI = False


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class PdfTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = ROOT
        cls.root.mkdir(parents=True, exist_ok=True)
        cls.source = cls.root / "원본-페이지.pdf"
        document = canvas.Canvas(str(cls.source), pagesize=(400, 500))
        for label in ("PAGE ONE", "PAGE TWO", "PAGE THREE"):
            document.drawString(40, 440, label)
            document.showPage()
        document.save()
        cls.second = cls.root / "두번째.pdf"
        document = canvas.Canvas(str(cls.second), pagesize=(400, 500))
        document.drawString(40, 440, "LAST PAGE")
        document.save()
        cls.table = cls.root / "문서-표.pdf"
        pdfmetrics.registerFont(TTFont("FixtureKorean", r"C:\Windows\Fonts\malgun.ttf"))
        document = canvas.Canvas(str(cls.table), pagesize=(600, 500))
        document.setFont("FixtureKorean", 15)
        document.drawString(40, 460, "FlipRing 한글 문서")
        for y, values in zip((380, 340, 300), (("Name", "Quantity", "Price"), ("Apple", "12", "300"), ("Pear", "25", "400"))):
            for x, value in zip((40, 220, 400), values):
                document.drawString(x, y, value)
        document.save()
        cls.image = cls.root / "한영-스캔.png"
        raster = Image.new("RGB", (1800, 1100), "white")
        draw = ImageDraw.Draw(raster)
        korean = ImageFont.truetype(r"C:\Windows\Fonts\malgun.ttf", 56)
        english = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 48)
        draw.text((90, 70), "한글 문서 인식 테스트", font=korean, fill="black")
        draw.text((90, 170), "FlipRing English OCR 2026", font=english, fill="black")
        for y, values in zip((340, 470, 600), (("Name", "Quantity", "Price"), ("Apple", "12", "300"), ("Pear", "25", "400"))):
            for x, value in zip((90, 700, 1200), values):
                draw.text((x, y), value, font=english, fill="black")
        raster.save(cls.image, dpi=(200, 200))
        cls.tiff = cls.root / "여러이미지.tiff"
        raster.save(cls.tiff, save_all=True, append_images=[Image.new("RGB", (600, 400), "gray")])
        cls.before = {path: digest(path) for path in (cls.source, cls.second, cls.table, cls.image, cls.tiff)}

    @classmethod
    def tearDownClass(cls):
        for path, before in cls.before.items():
            if digest(path) != before:
                raise AssertionError(f"원본 변경: {path}")

    def run_request(self, operation, paths=None, **options):
        return tools.execute(tools.Request(operation, [str(path) for path in (paths or [self.source])], **options))

    def test_01_ranges_safe_order(self):
        self.assertEqual(tools.page_indices("3,1-2,2", 3), [2, 0, 1])
        for value in ("0", "4", "2-1", "1,,2", "-1", "1-4", "a", "1;2"):
            with self.subTest(value=value), self.assertRaises(tools.PdfToolError):
                tools.page_indices(value, 3)
        with self.assertRaises(tools.PdfToolError):
            tools.page_indices("", 3, allow_empty=False)

    def test_02_merge_order(self):
        output = self.run_request("merge", [self.second, self.source])[0]
        reader = PdfReader(output)
        self.assertEqual(len(reader.pages), 4)
        self.assertIn("LAST PAGE", reader.pages[0].extract_text())
        self.assertIn("PAGE THREE", reader.pages[-1].extract_text())

    def test_03_split(self):
        outputs = self.run_request("split")
        self.assertEqual(len(outputs), 3)
        for number, output in enumerate(outputs):
            reader = PdfReader(output)
            self.assertEqual(len(reader.pages), 1)
            self.assertIn(("ONE", "TWO", "THREE")[number], reader.pages[0].extract_text())

    def test_04_extract_order_and_delete(self):
        output = self.run_request("extract", pages="3,1")[0]
        reader = PdfReader(output)
        self.assertEqual(len(reader.pages), 2)
        self.assertIn("THREE", reader.pages[0].extract_text())
        self.assertIn("ONE", reader.pages[1].extract_text())
        output = self.run_request("delete", pages="2")[0]
        self.assertEqual(len(PdfReader(output).pages), 2)
        self.assertIn("THREE", PdfReader(output).pages[-1].extract_text())
        with self.assertRaises(tools.PdfToolError) as error:
            self.run_request("delete", pages="1-3")
        self.assertEqual(error.exception.code, "all_pages_removed")

    def test_05_rotate_selected(self):
        output = self.run_request("rotate", pages="2", rotation=90)[0]
        pages = PdfReader(output).pages
        self.assertEqual([page.rotation for page in pages], [0, 90, 0])

    def test_06_image_pdf_multiple_frames(self):
        output = self.run_request("images_to_pdf", [self.image, self.tiff])[0]
        self.assertEqual(len(PdfReader(output).pages), 3)
        outputs = self.run_request("pdf_to_images", [output], pages="1,3", output_format="png", dpi=100)
        self.assertEqual(len(outputs), 2)
        for path in outputs:
            with Image.open(path) as image:
                image.verify()

    def test_07_lossless_compression_keeps_text(self):
        output = self.run_request("compress")[0]
        self.assertEqual(len(PdfReader(output).pages), 3)
        self.assertIn("PAGE TWO", PdfReader(output).pages[1].extract_text())

    def test_08_lossy_compression_renders(self):
        source = self.run_request("images_to_pdf", [self.image])[0]
        output = self.run_request("compress", [source], compression="images", dpi=80, quality=50)[0]
        document = tools.pdfium.PdfDocument(output)
        raster = self.root / "압축-시각확인.png"
        try:
            tools.render_page(document, 0, raster, 80)
        finally:
            document.close()
        with Image.open(raster) as image:
            self.assertGreater(image.width, 400)
        self.assertLess(Path(output).stat().st_size, Path(source).stat().st_size)

    def test_09_unique_names_and_source_hash(self):
        first = self.run_request("extract", pages="1")[0]
        second = self.run_request("extract", pages="1")[0]
        self.assertNotEqual(first, second)
        self.assertEqual(digest(self.source), self.before[self.source])
        self.assertEqual(Path(first).parent, self.source.parent)

    def test_10_encrypted_pdf_explicit_error(self):
        path = self.root / "암호.pdf"
        writer = PdfWriter()
        writer.append(str(self.source))
        writer.encrypt("secret")
        with path.open("wb") as stream:
            writer.write(stream)
        writer.close()
        with self.assertRaises(tools.PdfToolError) as error:
            self.run_request("rotate", [path])
        self.assertEqual(error.exception.code, "encrypted_pdf")

    def test_11_cancel_before_work(self):
        cancel = threading.Event()
        cancel.set()
        before = set(self.root.iterdir())
        with self.assertRaises(tools.Cancelled):
            tools.execute(tools.Request("split", [str(self.source)]), cancel)
        self.assertEqual(before, set(self.root.iterdir()))

    def test_12_coordinate_table_docx_xlsx(self):
        contents = tools.pdf_text_pages(self.table)
        rows = tools.structured_rows(contents[0])
        self.assertTrue(any(len(row) >= 3 and row[0] == "Apple" and row[1] == "12" for row in rows), rows)
        docx = self.run_request("recognize", [self.table], output_format="docx")[0]
        document = Document(docx)
        self.assertTrue(document.tables)
        self.assertTrue(any(len(row.cells) >= 3 and row.cells[0].text == "Apple" and row.cells[1].text == "12"
                            for table in document.tables for row in table.rows))
        xlsx = self.run_request("recognize", [self.table], output_format="xlsx")[0]
        workbook = load_workbook(xlsx, read_only=True)
        try:
            self.assertGreaterEqual(workbook.active.max_column, 3)
            self.assertTrue(any(row[0] == "Apple" and row[1] == "12" for row in workbook.active.values))
        finally:
            workbook.close()

    def test_13_formulas_are_literal_text(self):
        output = self.root / "수식-차단.xlsx"
        content = tools.PageContent("", [tools.Word("=1+1", 10, 10, 80, 25),
                                          tools.Word("safe", 180, 10, 220, 25)], 300, 300)
        tools.export_document([content], output, "xlsx")
        workbook = load_workbook(output)
        try:
            self.assertEqual(workbook.active["A1"].value, "=1+1")
            self.assertEqual(workbook.active["A1"].data_type, "s")
        finally:
            workbook.close()

    def test_14_external_process_timeout_and_temp_cleanup(self):
        work = self.root / "프로세스시간시험"
        work.mkdir(exist_ok=True)
        with self.assertRaises(tools.PdfToolError) as error:
            tools.run_ocr_process([sys.executable, "-c", "import time; time.sleep(30)"], work, threading.Event(), timeout=0.15)
        self.assertEqual(error.exception.code, "ocr_timeout")
        self.assertFalse(list(self.root.glob(".flipring-pdf-*")))

    def test_15_cancel_terminates_process_tree(self):
        work = self.root / "자식종료시험"
        work.mkdir(exist_ok=True)
        cancel = threading.Event()
        pid_file = work / "child.pid"
        pid_file.unlink(missing_ok=True)
        child_code = "import time; time.sleep(60)"
        parent_code = "import subprocess,sys,pathlib,time; p=subprocess.Popen([sys.executable,'-c'," + repr(child_code) + "]); pathlib.Path('child.pid').write_text(str(p.pid)); time.sleep(60)"

        def cancel_after_child():
            deadline = time.monotonic() + 10
            while not pid_file.exists() and time.monotonic() < deadline:
                time.sleep(0.02)
            cancel.set()

        thread = threading.Thread(target=cancel_after_child)
        thread.start()
        try:
            with self.assertRaises(tools.Cancelled):
                tools.run_ocr_process([sys.executable, "-c", parent_code], work, cancel)
        finally:
            thread.join(timeout=12)
        self.assertTrue(pid_file.is_file())
        pid = int(pid_file.read_text())
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = [ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32]
        kernel.OpenProcess.restype = ctypes.c_void_p
        kernel.WaitForSingleObject.argtypes = [ctypes.c_void_p, ctypes.c_uint32]
        kernel.CloseHandle.argtypes = [ctypes.c_void_p]
        handle = kernel.OpenProcess(0x100000, False, pid)
        if handle:
            try:
                self.assertEqual(kernel.WaitForSingleObject(handle, 5000), 0)
            finally:
                kernel.CloseHandle(handle)


class OcrTests(PdfTests):
    def test_16_real_bilingual_ocr_unicode_path(self):
        output = self.run_request("recognize", [self.image], output_format="txt", psm=6)[0]
        value = Path(output).read_text(encoding="utf-8")
        self.assertIn("English", value)
        self.assertIn("한글", value)
        self.assertIn("2026", value)

    def test_17_searchable_pdf_image_and_scanned_pdf(self):
        output = self.run_request("recognize", [self.image], output_format="pdf", psm=6)[0]
        self.assertIn("English", "".join(page.extract_text() for page in PdfReader(output).pages))
        self.assertIn("한글", "".join(page.extract_text() for page in PdfReader(output).pages))
        scanned = self.run_request("images_to_pdf", [self.image])[0]
        text = self.run_request("recognize", [scanned], output_format="txt", psm=6, dpi=200)[0]
        self.assertIn("English", Path(text).read_text(encoding="utf-8"))
        self.assertIn("한글", Path(text).read_text(encoding="utf-8"))
        document = tools.pdfium.PdfDocument(output)
        try:
            tools.render_page(document, 0, self.root / "검색PDF-시각확인.png", 100)
        finally:
            document.close()

    def test_18_real_ocr_docx_table_and_xlsx_cells(self):
        docx = self.run_request("recognize", [self.image], output_format="docx", psm=6)[0]
        document = Document(docx)
        self.assertTrue(document.tables)
        self.assertTrue(any(len(row.cells) >= 3 and "Apple" in row.cells[0].text and "12" in row.cells[1].text
                            for table in document.tables for row in table.rows))
        xlsx = self.run_request("recognize", [self.image], output_format="xlsx", psm=6)[0]
        workbook = load_workbook(xlsx, read_only=True)
        try:
            self.assertGreaterEqual(workbook.active.max_column, 3)
            self.assertTrue(any(row[0] and "Apple" in str(row[0]) and "12" in str(row[1])
                                for row in workbook.active.values))
        finally:
            workbook.close()


class PanelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.application = QApplication.instance() or QApplication([])
        cls.application.setQuitOnLastWindowClosed(False)

    def setUp(self):
        from pdf_panel import PdfToolsPanel
        from PySide6.QtCore import QThread, Slot

        class ObservingPanel(PdfToolsPanel):
            def __init__(self):
                self.affinity = []
                super().__init__("ko")

            @Slot(int, int, str)
            def _on_progress(self, done, total, label):
                self.affinity.append(("progress", QThread.currentThread() == self.thread()))
                super()._on_progress(done, total, label)

            @Slot(str, str, str)
            def _on_result(self, source, output, error):
                self.affinity.append(("result", QThread.currentThread() == self.thread()))
                super()._on_result(source, output, error)

            @Slot(object)
            def _on_error(self, error):
                self.affinity.append(("error", QThread.currentThread() == self.thread()))
                super()._on_error(error)

            @Slot(object)
            def _finish(self, future):
                self.affinity.append(("finish", QThread.currentThread() == self.thread()))
                super()._finish(future)

        self.panel = ObservingPanel()
        self.events = []
        self.panel.result_ready.connect(lambda source, output, error: self.events.append((source, output, error)))
        self.panel.set_files([str(ROOT / "원본-페이지.pdf"), str(ROOT / "두번째.pdf")])

    def tearDown(self):
        if self.panel.is_busy():
            self.panel.cancel()
            self.wait_finished()
        self.assertTrue(self.panel.shutdown(wait=True))
        self.panel.hide()
        self.panel.deleteLater()
        self.application.processEvents()

    def wait_finished(self):
        from PySide6.QtTest import QTest
        deadline = time.monotonic() + 10
        while self.panel.is_busy() and time.monotonic() < deadline:
            QTest.qWait(10)
        self.assertFalse(self.panel.is_busy())

    def test_19_panel_thread_result_language_and_restore(self):
        from PySide6.QtCore import QTimer
        ticks = []
        QTimer.singleShot(0, lambda: ticks.append(True))
        self.panel.start()
        self.assertTrue(self.panel.is_busy())
        self.assertFalse(self.panel.start_button.isEnabled())
        self.panel.set_language("en")
        self.wait_finished()
        self.assertTrue(ticks)
        self.assertTrue(self.panel.start_button.isEnabled())
        self.assertEqual(len(self.events), 1)
        self.assertTrue(Path(self.events[0][1]).is_file())
        self.assertEqual(self.events[0][2], "")
        self.assertIn("Done", self.panel.status.text())
        self.assertTrue(self.panel.affinity)
        self.assertTrue(all(matches for _, matches in self.panel.affinity), self.panel.affinity)

    def test_20_panel_error_restores_inputs(self):
        self.panel.operation.setCurrentIndex(self.panel.operation.findData("rotate"))
        self.panel.pages.setText("0")
        self.panel.start()
        self.wait_finished()
        self.assertTrue(self.panel.start_button.isEnabled())
        self.assertEqual(len(self.events), 1)
        self.assertEqual(self.events[0][1], "")
        self.assertIn("페이지", self.events[0][2])
        self.assertIn(("error", True), self.panel.affinity)

    def test_21_panel_cancel_is_responsive(self):
        from PySide6.QtTest import QTest

        def controlled_work(request, cancel, progress, result):
            # 실제 OCR 취소는 별도 검사했다. 여기서는 느린 작업을 제어하여 UI 응답성을 확인한다.
            while not cancel.wait(0.02):
                pass
            raise tools.Cancelled()

        with patch.object(tools, "execute", controlled_work):
            self.panel.start()
            QTest.qWait(30)
            self.assertTrue(self.panel.is_busy())
            self.panel.cancel()
            self.wait_finished()
        self.assertEqual(self.events, [])
        self.assertTrue(self.panel.start_button.isEnabled())
        self.assertIn("취소", self.panel.status.text())

    def test_22_panel_blocks_existing_engine_queue(self):
        self.panel.pending = 1
        self.panel.start()
        self.assertFalse(self.panel.is_busy())
        self.assertIn("다른 변환", self.panel.status.text())

    def test_23_worker_and_gui_thread_affinity(self):
        from PySide6.QtCore import QThread
        from pdf_panel import PdfToolsPanel
        observations = []
        main_python_thread = threading.get_ident()

        def observed_work(request, cancel, progress, result):
            observations.append(threading.get_ident() != main_python_thread)
            observations.append(QThread.currentThread() != self.application.thread())
            progress(1, 1, "thread test")
            result(request.paths[0], request.paths[0], "")

        for signature in ("_on_progress(int,int,QString)", "_on_result(QString,QString,QString)", "_on_error(PyObject)", "_finish(PyObject)"):
            self.assertGreaterEqual(PdfToolsPanel.staticMetaObject.indexOfSlot(signature), 0, signature)
        with patch.object(tools, "execute", observed_work):
            self.panel.start()
            self.wait_finished()
        self.assertEqual(observations, [True, True])
        self.assertIn(("progress", True), self.panel.affinity)
        self.assertIn(("result", True), self.panel.affinity)
        self.assertIn(("finish", True), self.panel.affinity)

    def test_24_repeat_completion_cancellation_and_shutdown(self):
        from PySide6.QtTest import QTest
        for _ in range(3):
            self.panel.start()
            self.wait_finished()
            self.assertIsNone(self.panel._future)
            self.assertTrue(self.panel.start_button.isEnabled())

        def controlled_work(request, cancel, progress, result):
            while not cancel.wait(0.02):
                pass
            raise tools.Cancelled()

        with patch.object(tools, "execute", controlled_work):
            self.panel.start()
            QTest.qWait(30)
            self.assertFalse(self.panel.shutdown())
            self.panel.cancel()
            self.wait_finished()
        self.panel.start()
        self.wait_finished()
        self.assertEqual(len(self.events), 4)
        self.assertTrue(self.panel.shutdown(wait=True))
        self.assertIsNone(self.panel._pool)

    def test_25_real_ocr_in_background_panel(self):
        if not OCR:
            self.skipTest("--ocr 인자가 필요한 실제 OCR 패널 시험")
        self.panel.set_files([str(ROOT / "한영-스캔.png")])
        self.panel.operation.setCurrentIndex(self.panel.operation.findData("recognize"))
        self.panel.output.setCurrentIndex(self.panel.output.findData("txt"))
        self.panel.psm.setCurrentIndex(self.panel.psm.findData(6))
        self.panel.start()
        self.wait_finished()
        self.assertEqual(len(self.events), 1)
        value = Path(self.events[0][1]).read_text(encoding="utf-8")
        self.assertIn("한글", value)
        self.assertIn("English", value)
        self.assertTrue(all(matches for _, matches in self.panel.affinity))

    def test_26_cancel_real_ocr_process_from_panel(self):
        if not OCR:
            self.skipTest("--ocr 인자가 필요한 실제 OCR 취소 시험")
        from PySide6.QtTest import QTest
        self.panel.set_files([str(ROOT / "한영-스캔.png")])
        self.panel.operation.setCurrentIndex(self.panel.operation.findData("recognize"))
        started = threading.Event()
        processes = []
        popen = tools.subprocess.Popen

        def observed_popen(*args, **kwargs):
            process = popen(*args, **kwargs)
            processes.append(process)
            started.set()
            return process

        with patch.object(tools.subprocess, "Popen", observed_popen):
            self.panel.start()
            deadline = time.monotonic() + 10
            while not started.is_set() and self.panel.is_busy() and time.monotonic() < deadline:
                QTest.qWait(5)
            self.assertTrue(started.is_set())
            self.panel.cancel()
            self.wait_finished()
        self.assertTrue(processes)
        self.assertTrue(all(process.poll() is not None for process in processes))
        self.assertEqual(self.events, [])
        self.assertIn("취소", self.panel.status.text())
        self.assertFalse(list(ROOT.glob(".flipring-pdf-*")))

    def test_27_second_input_failure_records_actual_source(self):
        invalid = ROOT / "두번째-파손.pdf"
        invalid.write_bytes(b"invalid pdf fixture")
        first = str(ROOT / "원본-페이지.pdf")
        self.panel.set_files([first, str(invalid)])
        self.panel.operation.setCurrentIndex(self.panel.operation.findData("rotate"))
        self.panel.start()
        self.wait_finished()
        self.assertEqual(len(self.events), 2)
        self.assertEqual(self.events[0][0], first)
        self.assertTrue(self.events[0][1])
        self.assertEqual(self.events[1][0], str(invalid.resolve()))
        self.assertEqual(self.events[1][1], "")
        self.assertTrue(self.events[1][2])
        self.assertTrue(self.panel.start_button.isEnabled())

    def test_28_merge_failure_has_source_and_global_error_has_none(self):
        invalid = ROOT / "합치기-파손.pdf"
        invalid.write_bytes(b"invalid merge pdf fixture")
        self.panel.set_files([str(ROOT / "원본-페이지.pdf"), str(invalid)])
        self.panel.start()
        self.wait_finished()
        self.assertEqual(len(self.events), 1)
        self.assertEqual(self.events[0][0], str(invalid.resolve()))
        self.events.clear()

        def global_failure(request, cancel, progress, result):
            raise tools.PdfToolError("operation")

        with patch.object(tools, "execute", global_failure):
            self.panel.start()
            self.wait_finished()
        self.assertEqual(len(self.events), 1)
        self.assertEqual(self.events[0][0], "")


def main():
    global ROOT, OCR, GUI
    parser = argparse.ArgumentParser()
    parser.add_argument("output")
    parser.add_argument("--ocr", action="store_true")
    parser.add_argument("--gui", action="store_true")
    args = parser.parse_args()
    ROOT = Path(args.output).resolve()
    OCR = args.ocr
    GUI = args.gui
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(OcrTests if OCR else PdfTests)
    if GUI:
        suite.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(PanelTests))
    result = unittest.TextTestRunner(verbosity=2, failfast=True).run(suite)
    report = {"tests": result.testsRun, "ocr_enabled": OCR, "gui_enabled": GUI,
              "passed": result.testsRun - len(result.failures) - len(result.errors) - len(result.skipped),
              "skipped": [{"test": str(test), "reason": reason} for test, reason in result.skipped],
              "failures": [{"test": str(test), "detail": detail} for test, detail in result.failures],
              "errors": [{"test": str(test), "detail": detail} for test, detail in result.errors],
              "scope": "생성한 PDF/이미지의 로컬 편집·원본해시·좌표표·외부프로세스 취소/시간제한. OCR활성시 실제 portable Tesseract 한영/검색PDF/DOCX표/XLSX셀 검증."}
    (ROOT / "PDF-OCR-검증.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in report.items() if key not in ("failures", "errors")}, ensure_ascii=False))
    return int(not result.wasSuccessful())


if __name__ == "__main__":
    sys.exit(main())
