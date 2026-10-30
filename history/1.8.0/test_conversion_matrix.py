"""지원표의 실제 변환을 독립 실행하고 결과와 원본을 검사한다."""
import argparse
import csv
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tarfile
import tempfile
import time
import traceback
import uuid
import zipfile
from datetime import datetime, timedelta, timezone

from PIL import Image
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
import engine

KST = timezone(timedelta(hours=9))
MARK = 'FlipRing matrix 2026'


def digest(path):
    if path.is_dir():
        return hashlib.sha256(b''.join(str(p.relative_to(path)).encode('utf-8') + p.read_bytes() for p in sorted(path.rglob('*')) if p.is_file())).hexdigest()
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fixture(root, rar_fixture=None):
    root.mkdir(parents=True, exist_ok=True)
    generated = {}
    errors = {}
    def make(ext, action):
        p = root / ('한글 입력 공백.' + ext)
        try:
            action(p)
            if not p.exists(): raise RuntimeError('시험 입력이 생성되지 않았습니다.')
            generated[ext] = str(p)
        except Exception:
            errors[ext] = traceback.format_exc()
    for ext in 'jpg png webp tiff avif bmp gif'.split():
        make(ext, lambda p, ext=ext: Image.new('RGB', (48, 36), (32, 122, 210)).save(p, format={'jpg':'JPEG','tiff':'TIFF'}.get(ext, ext.upper())))
    make('heic', lambda p: shutil.copyfile(Path(__file__).parent / 'test-fixtures/sample.heic', p))
    make('svg', lambda p: p.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="48" height="36"><rect width="48" height="36" fill="red"/></svg>', encoding='utf-8'))
    for ext, content in [('txt', MARK+' 한글 문서\n'), ('md', '# '+MARK+'\n\n한글 문서\n'), ('html', '<h1>'+MARK+'</h1><p>한글 문서</p>'), ('srt', '1\n00:00:00,000 --> 00:00:01,000\n한글 자막\n'), ('vtt', 'WEBVTT\n\n00:00.000 --> 00:01.000\n한글 자막\n')]:
        make(ext, lambda p, content=content: p.write_text(content, encoding='utf-8'))
    def doc(p):
        d=Document(); d.add_paragraph(MARK+' 한글 문서'); d.save(p)
    make('docx', doc)
    def presentation(p):
        d=Presentation(); d.slides.add_slide(d.slide_layouts[1]).shapes.title.text=MARK; d.save(p)
    make('pptx', presentation)
    def spreadsheet(p):
        w=Workbook(); w.active.append([MARK, '한글']); w.active.append(['값', 12]); w.save(p)
    make('xlsx', spreadsheet)
    for ext, delimiter in [('csv', ','), ('tsv', '\t')]:
        def sheet(p, delimiter=delimiter):
            with p.open('w', encoding='utf-8-sig', newline='') as f: csv.writer(f, delimiter=delimiter).writerows([[MARK,'한글'],['값','12']])
        make(ext, sheet)
    make('pdf', lambda p: engine.write_text_pdf((MARK+' 한글 PDF\n')*100, p))
    folder=root/'한글 폴더 공백'; folder.mkdir(); (folder/'원본.txt').write_text(MARK, encoding='utf-8'); (folder/'빈폴더').mkdir(); generated['folder']=str(folder)
    def converted(ext, origin, office_input=False):
        def action(p):
            source=Path(generated[origin])
            if office_input:
                # 입력 전용 형식의 자료 생성은 제품의 출력 지원표와 구분한다.
                with tempfile.TemporaryDirectory(prefix='office-fixture-', dir=root) as directory:
                    out=engine.office(source, ext, Path(directory))
                    shutil.copyfile(out,p)
            else:
                out=Path(engine.convert(source, ext))
                if out != p: shutil.copyfile(out,p); out.unlink()
        make(ext, action)
    for ext in ('doc','odt','rtf'): converted(ext,'docx',office_input=True)
    converted('epub','md')
    for ext in ('ppt','odp'): converted(ext,'pptx',office_input=True)
    for ext in ('xls','ods'): converted(ext,'xlsx',office_input=True)
    for ext in ('zip','tar','tar.gz'): converted(ext,'folder')
    rar=Path(rar_fixture) if rar_fixture is not None else None
    if rar is not None and rar.is_file(): make('rar', lambda p: shutil.copyfile(rar,p))
    else: errors['rar']='라이선스와 출처가 확인된 fixture.rar가 없습니다.'
    ff=engine.tool('ffmpeg')
    make('wav', lambda p: engine.run(ff,['-f','lavfi','-i','sine=frequency=440:duration=0.3',p]))
    for ext in 'mp3 m4a flac ogg opus aiff wma'.split(): converted(ext,'wav')
    make('mp4', lambda p: engine.run(ff,['-f','lavfi','-i','color=c=blue:s=64x48:d=0.3','-f','lavfi','-i','sine=frequency=440:duration=0.3','-c:v','libx264','-pix_fmt','yuv420p','-c:a','aac','-shortest',p]))
    for ext in 'mov mkv webm avi wmv'.split(): converted(ext,'mp4')
    return generated, errors


def validate(path, target, source_kind, source):
    if path.is_dir():
        items=list(path.iterdir())
        if len(items)<2: raise AssertionError('여러 페이지 이미지가 없습니다.')
        for p in items:
            with Image.open(p) as im: im.verify()
        return
    if not path.is_file(): raise AssertionError('출력 파일이 없습니다.')
    if path.stat().st_size == 0 and target not in ('txt','md'): raise AssertionError('빈 출력 파일입니다.')
    if target in ('jpg','png','webp','tiff','avif','bmp','gif'):
        with Image.open(path) as im: im.verify()
    elif target == 'pdf':
        import pypdfium2
        d=pypdfium2.PdfDocument(str(path))
        try:
            if len(d)<1: raise AssertionError('PDF 페이지가 없습니다.')
        finally: d.close()
    elif target == 'docx':
        if not Document(path).paragraphs: raise AssertionError('문서 문단이 없습니다.')
    elif target == 'xlsx':
        w=load_workbook(path, read_only=True); assert w.sheetnames; w.close()
    elif target in ('zip','epub','ods'):
        with zipfile.ZipFile(path) as z:
            if not z.namelist() or z.testzip() is not None: raise AssertionError('압축 내용이 손상됐습니다.')
            if target=='zip' and not any(not n.endswith('/') for n in z.namelist()): raise AssertionError('압축 파일 내용이 없습니다.')
    elif target in ('tar','tar.gz'):
        with tarfile.open(path) as t:
            if not any(m.isfile() for m in t.getmembers()): raise AssertionError('압축 파일 내용이 없습니다.')
    if target in ('zip','tar','tar.gz') and source_kind in ('폴더','압축'):
        def members(p, kind):
            if kind=='folder': return {p.name+'/'+q.relative_to(p).as_posix():q.read_bytes() for q in p.rglob('*') if q.is_file()}
            if kind=='zip':
                with zipfile.ZipFile(p) as z: return {n:z.read(n) for n in z.namelist() if not n.endswith('/')}
            if kind=='rar':
                import rarfile
                import zlib
                # 제품의 파일명 선택 어댑터와 독립적으로 전체 스트림을 검증한다.
                executable = Path(__file__).resolve().parent / 'vendor/7zip/7z.exe'
                result = subprocess.run([str(executable), 'e', '-so', '-bb0', '-p', '--', str(p)],
                                        stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                        timeout=30, creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
                if result.returncode: raise AssertionError('RAR 원문 해제 실패: ' + result.stderr.decode('utf-8', 'replace'))
                expected, offset = {}, 0
                with rarfile.RarFile(p) as z:
                    for entry in z.infolist():
                        if entry.is_dir(): continue
                        data = result.stdout[offset:offset + entry.file_size]
                        if len(data) != entry.file_size or zlib.crc32(data) != entry.CRC:
                            raise AssertionError('RAR 원문 순서·크기·CRC 검증 실패')
                        name = entry.filename.replace('\\', '/')
                        if name in expected: raise AssertionError('RAR 원문 파일 이름 중복')
                        expected[name] = data
                        offset += entry.file_size
                if offset != len(result.stdout): raise AssertionError('RAR 원문에 예상하지 않은 후행 데이터가 있습니다.')
                return expected
            with tarfile.open(p) as t: return {x.name:t.extractfile(x).read() for x in t.getmembers() if x.isfile()}
        actual=members(path,target); expected=members(source,engine.classify(source)[1])
        if actual != expected: raise AssertionError('압축 구성 파일명 또는 내용이 원본과 다릅니다.')
    elif target in ('srt','vtt'):
        value=path.read_text(encoding='utf-8-sig'); assert '-->' in value and '한글 자막' in value
    elif target in ('csv','tsv'):
        with path.open(encoding='utf-8-sig',newline='') as f: assert list(csv.reader(f,delimiter='\t' if target=='tsv' else ','))
    elif target in ('mp3','m4a','wav','flac','ogg','opus','aiff','wma','mp4','mov','mkv','webm','avi','wmv'):
        ff=engine.tool('ffmpeg')
        p=subprocess.run([ff,'-v','error','-i',str(path),'-f','null','-'],capture_output=True,timeout=30)
        if p.returncode: raise RuntimeError(p.stderr.decode('utf-8','replace'))
    elif target in ('txt','md','html'):
        value=path.read_text(encoding='utf-8-sig')
        if source_kind != 'PDF' and not value.strip(): raise AssertionError('문서 내용이 비어 있습니다.')


def worker(source, target):
    before=digest(source)
    output=Path(engine.convert(source,target))
    validate(output,target,engine.classify(source)[0],source)
    if digest(source)!=before: raise AssertionError('원본 해시가 변경됐습니다.')
    return {'output':output.name,'source_sha256':before,'output_sha256':digest(output)}


def main():
    p=argparse.ArgumentParser(); p.add_argument('--root',type=Path); p.add_argument('--worker',nargs=2); p.add_argument('--fixtures',type=Path); p.add_argument('--rar-fixture',type=Path); p.add_argument('--timeout',type=int,default=35); p.add_argument('--total-timeout',type=int,default=1200); a=p.parse_args()
    if a.fixtures:
        print(json.dumps(fixture(a.fixtures,a.rar_fixture),ensure_ascii=False)); return 0
    if a.worker:
        try: print(json.dumps(worker(Path(a.worker[0]),a.worker[1]),ensure_ascii=False)); return 0
        except Exception: traceback.print_exc(); return 1
    if a.root is None: p.error('--root가 필요합니다')
    a.root = a.root.resolve()
    a.root.mkdir(parents=True,exist_ok=False)
    start=time.monotonic()
    fixture_cmd=[sys.executable,'-B','-X','utf8',str(Path(__file__).resolve()),'--fixtures',str(a.root/'fixtures')]
    if a.rar_fixture is not None: fixture_cmd.extend(['--rar-fixture',str(a.rar_fixture.resolve())])
    process=subprocess.Popen(fixture_cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,cwd=Path(__file__).parent)
    try:
        out,err=process.communicate(timeout=min(300,a.total_timeout))
        (a.root/'fixture.stdout.bin').write_bytes(out)
        (a.root/'fixture.stderr.bin').write_bytes(err)
        if process.returncode: raise RuntimeError(err.decode('utf-8','replace'))
        inputs,fixture_errors=json.loads(out.decode('utf-8'))
    except subprocess.TimeoutExpired:
        engine.terminate_tree(process)
        raise RuntimeError('시험 입력 생성 시간 제한 초과')
    expected=[(kind,src,dst) for kind,(srcs,dsts) in engine.TABLE.items() for src in srcs.split() for dst in dsts.split() if src!=dst]
    records=[]; tools={n:engine.tool(n) for n in ('ffmpeg','pandoc','soffice')}
    tool_info={k:({'name':Path(v).name,'sha256':digest(Path(v))} if v else None) for k,v in tools.items()}
    for kind,src,dst in expected:
        began=time.monotonic(); entry={'kind':kind,'source':src,'target':dst,'time_kst':datetime.now(KST).isoformat(),'tools':tool_info}
        if src not in inputs:
            entry.update(status='blocked_fixture',error=fixture_errors.get(src,'fixture 없음'))
        elif time.monotonic()-start >= a.total_timeout:
            entry.update(status='blocked_dependency',error='전체 시험 시간 제한')
        else:
            source=Path(inputs[src]); before=digest(source); entry['source_sha256']=before
            missing=engine.requirement(source,dst)
            cmd=[sys.executable,'-B','-X','utf8',str(Path(__file__).resolve()),'--worker',str(source),dst]
            entry['command']=cmd; entry['cwd']=str(Path(__file__).parent)
            if missing: entry.update(status='blocked_dependency',error=missing+' 도구가 없습니다.')
            else:
                try:
                    proc=subprocess.Popen(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,cwd=Path(__file__).parent)
                    try:
                        out,err=proc.communicate(timeout=min(a.timeout,max(1,a.total_timeout-(time.monotonic()-start))))
                    except subprocess.TimeoutExpired:
                        engine.terminate_tree(proc)
                        raise
                    entry.update(exit_code=proc.returncode,stdout=out.decode('utf-8','replace'),stderr=err.decode('utf-8','replace'),status='passed' if proc.returncode==0 else 'failed')
                    if digest(source)!=before: entry.update(status='failed',error='원본 SHA 변경')
                except subprocess.TimeoutExpired as e:
                    entry.update(status='failed',error='개별 실행 제한 시간 초과',stdout=(e.stdout or b'').decode('utf-8','replace') if isinstance(e.stdout,bytes) else e.stdout,stderr=(e.stderr or b'').decode('utf-8','replace') if isinstance(e.stderr,bytes) else e.stderr)
        entry['duration_seconds']=round(time.monotonic()-began,3); records.append(entry)
        (a.root/'progress.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps({k:entry.get(k) for k in ('kind','source','target','status','duration_seconds')},ensure_ascii=False),flush=True)
    extras=[]
    def extra(name, action):
        try:
            action(); extras.append({'name':name,'status':'passed'})
        except Exception:
            extras.append({'name':name,'status':'failed','error':traceback.format_exc()})
    if all(k in inputs for k in ('srt','txt')):
        def intersection():
            assert engine.targets([inputs['srt'],inputs['txt']]) == []
        extra('지원 교집합',intersection)
        def collision():
            source=Path(inputs['srt']); before=digest(source)
            first=Path(engine.convert(source,'vtt')); first_hash=digest(first)
            second=Path(engine.convert(source,'vtt'))
            assert first != second and digest(first)==first_hash and digest(source)==before
        extra('파일명 충돌과 원본 보존',collision)
    def traversal():
        source=a.root/'경로 탈출.zip'
        with zipfile.ZipFile(source,'w') as z: z.writestr('../outside.txt','거부')
        before=digest(source)
        try: engine.convert(source,'tar')
        except ValueError: pass
        else: raise AssertionError('상위 경로를 거부하지 않았습니다.')
        assert digest(source)==before and not (a.root.parent/'outside.txt').exists()
    extra('압축 상위 경로 거부',traversal)
    keys=[(r['kind'],r['source'],r['target']) for r in records]
    summary={'run_id':str(uuid.uuid4()),'expected':len(expected),'recorded':len(records),'executed':sum(r['status'] in ('passed','failed') for r in records),'missing':[x for x in expected if x not in keys],'duplicate':len(keys)-len(set(keys)),'counts':{s:sum(r['status']==s for r in records) for s in ('passed','failed','blocked_fixture','blocked_dependency')},'fixture_errors':fixture_errors,'extras':extras,'duration_seconds':round(time.monotonic()-start,3),'results':records}
    (a.root/'report.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    return int(summary['counts']['failed']>0 or summary['counts']['blocked_fixture']>0 or summary['counts']['blocked_dependency']>0 or bool(summary['missing']) or bool(summary['duplicate']) or any(x['status']=='failed' for x in extras))


if __name__=='__main__': sys.exit(main())
